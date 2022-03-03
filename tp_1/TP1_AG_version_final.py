#Version final
from random import randint, random
from openpyxl import  load_workbook 
from openpyxl.styles import Font, Color
from openpyxl.chart import Reference, LineChart, Series

ciclos = 20
tipoSeleccion = 1 #1=RULETA 2=TORNEO
elite = False #True/False

cantGenes = 30
cantCromosomas = 10
probCross = 0.75
probMutac = 0.05 
pobEntero = [0] * cantCromosomas
valoresFuncObj = [0] * cantCromosomas
valoresFuncFit = [0] * cantCromosomas
minimosFO = [] #Contiene el valor del minimo de la FO de cada corrida (20 generaciones)
maximosFO = []
promediosFO = []
cromosomasMax = [] #Contiene el valor binario del mayor cromosoma de cada generacion 
cromosomasMaxDecimal = []

#Generacion de la poblacion inicial en binario (gen a gen)
def generaPobBin():
    pobInicial = []
    for i in range(cantCromosomas):
        pobInicial.append([])
        for j in range(cantGenes):
            pobInicial[i].append(None) #Creo la estructura de la matriz
            pobInicial[i][j]=(randint(0,1)) #Lleno cada posicion con 0 o 1
    return pobInicial

#Pasaje de un cromosoma recibido en binario a decimal. Recibe un cromosoma en binario y retorna su correspondiente valor decimal.
def deBinaEnt(cromosoma):
    exponentes = []
    for j in range(cantGenes-1, -1, -1): #Itero de atras hacia adelante, para obtener los exponentes segun su posicion
            exponentes.append(j)
    acum = 0
    for j in range(cantGenes):
        acum = acum + (cromosoma[j]*(2**exponentes[j])) 
    entero = acum  
    return(entero)

#Calculo de la funcion objetivo para un cromosoma decimal recibido
def funcObjetivo(nro):
    x = (nro/((2**30) - 1))**2
    return(x)

#Operador de CROSSOVER 
def crossover(cromo1, cromo2):
    probabilidad = random()
    nuevo_cromo1 = cromo1.copy() #Por si no se da el crossover, que devuelva los padres
    nuevo_cromo2 = cromo2.copy()
    if probabilidad <= probCross: #Verifico que se cumpla la prob de crossover
        corte = randint(0, cantGenes-1)  #Punto de corte del crossover (cantGenes-1 para que no se salga de index)
        nuevo_cromo1 = cromo1[:corte] + cromo2[corte:]
        nuevo_cromo2 = cromo2[:corte] + cromo1[corte:]
    return nuevo_cromo1,nuevo_cromo2 

#Operador de MUTACION
def mutacion(cromo):
    probabilidad = random()
    if probabilidad <= probMutac: #Verifico que se cumpla la prob de mutacion
        gen = randint(0, cantGenes-1) #(cantGenes-1 para que no se salga de index)
        if cromo[gen] == 0: #Ejecuto la mutacion modificando directamente el cromosoma
            cromo[gen] = 1
        else:
            cromo[gen] = 0
    return cromo

#Armado de ruleta. 
#Recibe la lista de los fitness de cada cromosoma, y genera un nuevo arreglo (la ruleta) con el fitness acumulado de cada cromosoma, el cual se referenciara segun su index.
def armar_ruleta(list_fitnesses_prob):
    sum_fitnesses = 0
    for fit in list_fitnesses_prob:
        sum_fitnesses += fit

    list_fitnesses_prob_acum = [] #Ruleta
    prob_acum = 0.0

    for prob_fit in list_fitnesses_prob:
        prob_acum += prob_fit
        list_fitnesses_prob_acum.append(prob_acum)

    return list_fitnesses_prob_acum


#Tirado de ruleta. 
#Recibe el arreglo de fitness acumulados (ruleta), genera un numero random (gira la ruleta), y devuelve el indice correspondiente al cromosoma elegido en la ruleta (segun su fitness acumulado).
def tirar_ruleta(list_fitnesses_prob_acum):
    random_value = random()

    for index, prob in enumerate(list_fitnesses_prob_acum): #Se guarda el indice y el contenido de list_fitnesses_prob_acum
        if random_value <= prob:
            return index

#Metodo de seleccion por TORNEO.
#Se eligen participantes del torneo para encontrar un ganador segun su fitness.El ganador sera seleccionado como padre
def torneo(lista_fitness):
    tam_torneo = randint(2,cantCromosomas)
    participantes = [] #Se llena con fitness de cada cromo participante
    indices_participantes = [] 
    cant_participantes = 0
    while (cant_participantes < tam_torneo): #Busco aprticipantes hasta llenar el torneo
        indice_fitness_participante = randint(0, cantCromosomas-1) #Referencio dentro del index del arreglo, que arranca en 0 y termina en cantCromo-1
        if indice_fitness_participante not in indices_participantes: #Verifico que no sea participante (elegido antes)
            participantes.append( lista_fitness[indice_fitness_participante] )
            indices_participantes.append(indice_fitness_participante)
            cant_participantes = cant_participantes + 1

    #Una vez que tengo todos los participantes busco el ganador del torneo:
    ganador = participantes[0]
    indice_ganador = indices_participantes[0]
    for i in range(tam_torneo-1):
        if participantes[i+1] > ganador: #Comparo segun fitness
            ganador = participantes[i+1]
            indice_ganador = indices_participantes[i+1]

    return indice_ganador #Retorna el indice al cromosoma ganador

#Aplicacion de ELITISMO. Devuelve el indice de los 2 mejores cromosomas (elites) de la poblacion, segun los valores de fitness del ciclo.
def elitismo():
    mayorFit = valoresFuncFit[0]
    indiceElit1 = 0
    mayorFit2 = 0
    for i in range(1, cantCromosomas): #arranco en 1 porque el 0 ya fue considerado
        if(valoresFuncFit[i]>=mayorFit):
            mayorFit2 = mayorFit
            indiceElit2 = indiceElit1
            mayorFit = valoresFuncFit[i]
            indiceElit1 = i
        elif(valoresFuncFit[i]>mayorFit2):
            mayorFit2 = valoresFuncFit[i]
            indiceElit2 = i        
    return indiceElit1, indiceElit2

#Aplicacion del algoritmo genetico (orden de sucesos y llamadas a funciones)
def aplicarAG(poblacion, respuesta): 
    for i in range(cantCromosomas): 
        pobEntero[i] = deBinaEnt(poblacion[i]) #Pasaje a entero de la poblacion binaria generada
        valoresFuncObj[i] = funcObjetivo(pobEntero[i]) #Valores de la funcion objetivo  
    
    #Calculos de la funcion objetivo
    sumaObj = 0
    mayorObj = valoresFuncObj[0]
    menorObj = valoresFuncObj[0]
    indiceMayorObj = 0
    for i in range(cantCromosomas):
        sumaObj = sumaObj + valoresFuncObj[i]
        if (valoresFuncObj[i]>=mayorObj): 
            mayorObj = valoresFuncObj[i]
            indiceMayorObj = i
        if (valoresFuncObj[i]<=menorObj): 
            menorObj = valoresFuncObj[i]
    promObj = sumaObj/cantCromosomas
    cromosomaMayor = poblacion[indiceMayorObj].copy()

    #Guardo los valores de cada generacion para la tabla
    minimosFO.append(menorObj)
    maximosFO.append(mayorObj)
    promediosFO.append(promObj)
    cromosomasMax.append(cromosomaMayor)
    cromosomasMaxDecimal.append(deBinaEnt(cromosomaMayor))

    #Calculo de funcion fitness
    for i in range(cantCromosomas):
        valoresFuncFit[i] = (valoresFuncObj[i]/sumaObj) #valores de la funcion fitness

    #comienza la seleccion de padres
    padres_selec = []

    if(respuesta == 1): #RULETA
        #armado de ruleta
        list_prob_fit_acum = armar_ruleta(valoresFuncFit)

        if elite: #Con elitismo
            elit1, elit2 = elitismo() #Indice de los cromo elite
            for i in range(cantCromosomas-2): #Elijo 8 padres
                index = tirar_ruleta(list_prob_fit_acum) #Devuelve el indice de 1 padre
                padres_selec.append(poblacion[index])
        else: #Sin elitismo
            for i in range(cantCromosomas): #Elijo 10 padres
                index = tirar_ruleta(list_prob_fit_acum)
                padres_selec.append(poblacion[index])

    if(respuesta == 2): #TORNEO
        #seleccion de padres
        if elite: #Con elitismo (elijo 8 padres) 
            elit1, elit2 = elitismo()
            cantPadres = cantCromosomas-2
        else: #Sin elitismo 
            cantPadres = cantCromosomas

        for i in range(cantPadres):
            #comienzo de torneo
            indice_ganador = torneo(valoresFuncFit)
            padres_selec.append(poblacion[indice_ganador])

    #aplicar crossover a cada cromosoma de la poblacion
    pob_cross = []
    #Crossover de los padres seleccionados, teniendo en cuenta si se aplica elitismo o no.
    if elite:
        ultIndice = cantCromosomas-3
    else:
        ultIndice = cantCromosomas-1
    for i in range(0, ultIndice, 2):
        crom1, crom2 = crossover(padres_selec[i], padres_selec[i+1])
        pob_cross.append(crom1)
        pob_cross.append(crom2)  

    #generar nueva poblacion aplicando mutacion a los cromosomas con crossover ya calculado
    nueva_poblacion = [] #Con elitismo: 8 hijos. Sin elitismo: 10 hijos
    #aplico mutacion
    if elite: #Con elitismo
        nueva_poblacion.append(poblacion[elit1]) #Agrego los dos elite a la nueva poblacion
        nueva_poblacion.append(poblacion[elit2])
        for i in range(cantCromosomas-2): #Mutacion a los 8 hijos
            nueva_poblacion.append(mutacion(pob_cross[i]))   
    else: #Sin elitismo
        for i in range(cantCromosomas): #Mutacion a todos (10 hijos)
            nueva_poblacion.append(mutacion(pob_cross[i]))
    
    #Conversion a decimal de la nueva poblacion binaria
    nueva_poblacion_entero = []
    for i in nueva_poblacion:
        nueva_poblacion_entero.append(deBinaEnt(i))

    return nueva_poblacion   

#Creacion de tablas y graficos de excel, donde se expondran los resultados del programa.
def tablaExcel(nombreArchivo):
    wb = load_workbook(filename=nombreArchivo) #para arbir uno existente, con los colores preseteados
    tabla = wb.active
    
    #Busuqeda del indice del cromosoma con mayor funcion objetivo de todas las generaciones.
    indiceMayorTotal = 0
    mayorTotal = maximosFO[0]
    for i in range(1, ciclos):
        if (maximosFO[i]>mayorTotal):
            mayorTotal = maximosFO[i]
            indiceMayorTotal = i 
    tabla["A1"] = "Maximo cromosoma: "
    tabla["C1"] = cromosomasMaxDecimal[indiceMayorTotal]
    tabla['A1'].font = Font(bold=True, size=12, color="A72D13") #Estilo visual
    tabla['C1'].font = Font(bold=True, size=12, color="A72D13")

    tabla.append(["Generacion", "Minimo FO", "Maximo FO", "ValorDecimal", "Cromosoma Binario", "Promedio FO"]) #titulos
    tabla['A2'].font = Font(bold=True) #Pongo los titulos en negrita. A modo visual
    tabla['B2'].font = Font(bold=True)
    tabla['C2'].font = Font(bold=True)
    tabla['D2'].font = Font(bold=True)
    tabla['E2'].font = Font(bold=True)
    tabla['F2'].font = Font(bold=True)
    for i in range(ciclos): #Datos de la corrida
        tabla.append([i+1, minimosFO[i], maximosFO[i], cromosomasMaxDecimal[i], str(cromosomasMax[i]), promediosFO[i]])
    
    #Generacion de grafico de lineas con minimos, maximos y promedios.
    grafica = LineChart()
    grafica.title = "Grafico de FO" 
    grafica.style = 10
    grafica.y_axis.title = "Valores FO"
    grafica.x_axis.title = "Ciclos"
    grafica.y_axis.scaling.min = 0 #Limites del eje y
    grafica.y_axis.scaling.max = 1
    grafica.x_axis.scaling.min = 1 #Limites del eje x
    grafica.x_axis.scaling.max = ciclos 
    data = Reference(tabla, min_col=2, min_row=2, max_col=3, max_row=ciclos+2) #Rango que se grafica MaximoFO, MinimoFO
    data2 = Reference(tabla, min_col=6, min_row=2, max_row=ciclos+2)#Columna de Promedio FO
    grafica.add_data(data, titles_from_data=True)
    grafica.add_data(data2, titles_from_data=True)
    finTabla = "B" + str(ciclos+5) #Para ubicar el grafico donde termina la tabla (dependendiendo la cant de ciclos)
    tabla.add_chart(grafica, finTabla) #agregado y ubicacion del grafico
    
    wb.save(nombreArchivo)
    
#Programa principal
poblacion = generaPobBin() #Genero poblacion inicial en binario. Unico en el programa

pob_nueva = aplicarAG(poblacion, tipoSeleccion)
for i in range(ciclos-1):
    pob_nueva = aplicarAG(pob_nueva, tipoSeleccion)

tablaExcel('tabla.xlsx') 

print('-- FIN DE LA CORRIDA --')
